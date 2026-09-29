"""
Script export to downloadable documents (TXT and DOCX).
"""

from __future__ import annotations

import io
from enum import StrEnum

from docx import Document
from docx.shared import Pt

TITLE_FONT_SIZE = Pt(16)


class ExportFormat(StrEnum):
    """Supported export formats."""

    TXT = "txt"
    DOCX = "docx"


#: MIME type for each export format.
MEDIA_TYPES: dict[ExportFormat, str] = {
    ExportFormat.TXT: "text/plain; charset=utf-8",
    ExportFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def export_script(topic: str, script: str, fmt: ExportFormat) -> bytes:
    """Render the script in the requested format and return the file bytes."""
    if fmt is ExportFormat.DOCX:
        return script_to_docx(topic, script)
    return script_to_txt(topic, script)


def script_to_txt(topic: str, script: str) -> bytes:
    """Plain-text export: topic header followed by the script."""
    return f"TOPIC: {topic}\n\nSCRIPT:\n{script}".encode("utf-8")


def script_to_docx(topic: str, script: str) -> bytes:
    """Word export: bold title, a ``SCRIPT:`` heading, then one paragraph per line."""
    document = Document()

    title = document.add_paragraph().add_run(topic)
    title.bold = True
    title.font.size = TITLE_FONT_SIZE

    document.add_paragraph("")
    document.add_paragraph().add_run("SCRIPT:").bold = True

    for line in script.split("\n"):
        document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
