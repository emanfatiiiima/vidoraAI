"""
Project package export (Step 8 "Download").

Bundles everything produced by the pipeline into a single ZIP file::

    <topic>.zip
    ├── project_details.xlsx   # Overview, Scenes (text + prompts) and Files sheets
    ├── script.txt             # full narration script + summary
    ├── images/img1.png ...    # generated scene images
    ├── videos/vid1.mp4 ...    # generated scene clips
    ├── audio/narration.mp3    # TTS narration
    └── final_video.mp4        # compiled video

Media is downloaded from its URL (or decoded from ``data:`` URIs). A file that
can't be fetched doesn't fail the whole export; it's marked as failed in the
"Files" sheet instead.
"""

from __future__ import annotations

import asyncio
import base64
import io
import logging
import mimetypes
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from urllib.parse import urlsplit

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from core.utils import fetch_bytes

logger = logging.getLogger(__name__)

# Downloads run in parallel across hosts but one at a time per host: free hosts
# (e.g. Pollinations) answer parallel requests with 402/429.
MAX_CONCURRENT_PER_HOST = 1
# Hard limit per file (including retries) so the export can never hang.
ASSET_TIMEOUT_SECONDS = 90.0
DETAILS_FILENAME = "project_details.xlsx"
SCRIPT_FILENAME = "script.txt"

# Extension used when the server doesn't tell us the content type.
_DEFAULT_EXTENSION = {"image": ".png", "video": ".mp4", "audio": ".mp3"}
# mimetypes returns odd defaults for some common types; pin the usual ones.
_PREFERRED_EXTENSION = {
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/webp": ".webp",
    "video/mp4": ".mp4",
    "audio/mpeg": ".mp3",
}
_DATA_URI_RE = re.compile(r"^data:(?P<mime>[\w/+.-]+)?(?:;[\w=-]+)*;base64,(?P<data>.*)$", re.DOTALL)


# ---------------------------------------------------------------------- #
# Input data
# ---------------------------------------------------------------------- #
@dataclass(slots=True)
class ProjectScene:
    """One scene as produced by Steps 5-6."""

    text: str = ""
    prompt: str = ""
    image_url: str | None = None
    video_url: str | None = None


@dataclass(slots=True)
class ProjectData:
    """Everything the user produced in the pipeline."""

    niche: str = ""
    style: str = ""
    duration: str = ""
    aspect_ratio: str = ""
    scene_count: int | None = None
    topic: str = ""
    script: str = ""
    summary: str = ""
    audio_url: str | None = None
    final_video_url: str | None = None
    scenes: list[ProjectScene] = field(default_factory=list)


@dataclass(slots=True)
class _Asset:
    """A media file to put in the ZIP (``name`` has no extension yet)."""

    kind: str  # "image" | "video" | "audio"
    name: str  # e.g. "images/img1"
    url: str
    archive_path: str = ""
    error: str | None = None


# ---------------------------------------------------------------------- #
# Exporter
# ---------------------------------------------------------------------- #
class ProjectExporter:
    """Builds the downloadable ZIP package for a finished project."""

    def __init__(self, http: httpx.AsyncClient) -> None:
        self._http = http

    async def build_zip(self, project: ProjectData) -> bytes:
        """Download all media and return the ZIP archive as bytes."""
        assets = _collect_assets(project)
        downloaded = await self._download_all(assets)

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for asset in assets:
                content = downloaded.get(asset.name)
                if content is not None:
                    # Media is already compressed; storing it is faster and just as small.
                    archive.writestr(asset.archive_path, content, compress_type=zipfile.ZIP_STORED)

            archive.writestr(SCRIPT_FILENAME, _script_text(project))
            archive.writestr(DETAILS_FILENAME, _details_workbook(project, assets))

        return buffer.getvalue()

    async def _download_all(self, assets: list[_Asset]) -> dict[str, bytes]:
        """Fetch every asset. The same URL is only downloaded once."""
        host_locks: defaultdict[str, asyncio.Semaphore] = defaultdict(
            lambda: asyncio.Semaphore(MAX_CONCURRENT_PER_HOST)
        )
        cache: dict[str, asyncio.Task[tuple[bytes, str | None]]] = {}

        async def fetch(url: str) -> tuple[bytes, str | None]:
            async with host_locks[urlsplit(url).netloc]:
                return await asyncio.wait_for(self._fetch(url), ASSET_TIMEOUT_SECONDS)

        for asset in assets:
            if asset.url not in cache:
                cache[asset.url] = asyncio.create_task(fetch(asset.url))

        results: dict[str, bytes] = {}
        for asset in assets:
            try:
                content, content_type = await cache[asset.url]
            except Exception as exc:  # noqa: BLE001 - one bad file must not break the export
                logger.warning("Could not download %s (%s): %s", asset.name, asset.url[:80], exc)
                asset.error = _short_error(exc)
                asset.archive_path = asset.name + _DEFAULT_EXTENSION[asset.kind]
                continue
            asset.archive_path = asset.name + _extension_for(content_type, asset.kind)
            results[asset.name] = content
        return results

    async def _fetch(self, url: str) -> tuple[bytes, str | None]:
        """Return ``(content, content_type)`` for an http(s) URL or ``data:`` URI."""
        if url.startswith("data:"):
            match = _DATA_URI_RE.match(url)
            if not match:
                raise ValueError("Unsupported data URI")
            return base64.b64decode(match.group("data")), match.group("mime")

        return await fetch_bytes(self._http, url)


# ---------------------------------------------------------------------- #
# Helpers
# ---------------------------------------------------------------------- #
def _collect_assets(project: ProjectData) -> list[_Asset]:
    """List every media file of the project with its target name in the ZIP."""
    assets: list[_Asset] = []
    for i, scene in enumerate(project.scenes, start=1):
        if scene.image_url:
            assets.append(_Asset("image", f"images/img{i}", scene.image_url))
    for i, scene in enumerate(project.scenes, start=1):
        if scene.video_url:
            assets.append(_Asset("video", f"videos/vid{i}", scene.video_url))
    if project.audio_url:
        assets.append(_Asset("audio", "audio/narration", project.audio_url))
    if project.final_video_url:
        assets.append(_Asset("video", "final_video", project.final_video_url))
    return assets


def _short_error(exc: Exception) -> str:
    """One-line, human-readable reason for a failed download (for the Files sheet)."""
    if isinstance(exc, httpx.HTTPStatusError):
        return f"HTTP {exc.response.status_code} {exc.response.reason_phrase}"
    if isinstance(exc, (TimeoutError, httpx.TimeoutException)):
        return "timed out"
    return str(exc).splitlines()[0] if str(exc) else type(exc).__name__


def _extension_for(content_type: str | None, kind: str) -> str:
    """File extension for a MIME type, falling back to a sensible default per kind."""
    if content_type:
        if content_type in _PREFERRED_EXTENSION:
            return _PREFERRED_EXTENSION[content_type]
        guessed = mimetypes.guess_extension(content_type)
        if guessed:
            return guessed
    return _DEFAULT_EXTENSION[kind]


def _script_text(project: ProjectData) -> str:
    """Plain-text version of the script with its summary."""
    return (
        f"TOPIC: {project.topic}\n\n"
        f"SUMMARY:\n{project.summary or '-'}\n\n"
        f"SCRIPT:\n{project.script}\n"
    )


def _details_workbook(project: ProjectData, assets: list[_Asset]) -> bytes:
    """Excel workbook with three sheets: Overview, Scenes and Files."""
    workbook = Workbook()

    overview = workbook.active
    overview.title = "Overview"
    _write_table(
        overview,
        ["Field", "Value"],
        [
            ["Niche", project.niche],
            ["Topic", project.topic],
            ["Visual style", project.style],
            ["Duration", project.duration],
            ["Aspect ratio", project.aspect_ratio],
            ["Scene count", project.scene_count if project.scene_count is not None else len(project.scenes)],
            ["Summary", project.summary],
            ["Script", project.script],
            ["Narration audio URL", project.audio_url or ""],
            ["Final video URL", project.final_video_url or ""],
        ],
        widths=[22, 100],
    )

    image_files = {a.name: a.archive_path for a in assets if a.kind == "image"}
    video_files = {a.name: a.archive_path for a in assets if a.kind == "video"}
    _write_table(
        workbook.create_sheet("Scenes"),
        ["Scene", "Narration text", "Image prompt", "Image file", "Video file"],
        [
            [
                i,
                scene.text,
                scene.prompt,
                image_files.get(f"images/img{i}", ""),
                video_files.get(f"videos/vid{i}", ""),
            ]
            for i, scene in enumerate(project.scenes, start=1)
        ],
        widths=[8, 60, 80, 18, 18],
    )

    _write_table(
        workbook.create_sheet("Files"),
        ["File", "Status", "Source URL"],
        [
            [
                a.archive_path,
                f"missing: {a.error}" if a.error else "included",
                # data: URIs can be megabytes long; don't dump them into a cell.
                "(embedded image data)" if a.url.startswith("data:") else a.url,
            ]
            for a in assets
        ],
        widths=[24, 40, 100],
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_table(sheet: Worksheet, headers: list[str], rows: list[list], widths: list[int]) -> None:
    """Write a header row plus data rows with wrapped text and fixed column widths."""
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column, width in zip("ABCDEFGHIJ", widths, strict=False):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
