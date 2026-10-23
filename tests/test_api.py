"""HTTP-level tests for ``fastapi_backend`` (using fake services)."""

from __future__ import annotations


def test_health(client) -> None:
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok", "llmProviders": ["fake"]}


def test_generate_returns_text_and_provider(client, fake_llm) -> None:
    fake_llm.answers = ["hello"]
    response = client.post(
        "/api/generate",
        json={"prompt": "hi", "responseMimeType": "application/json"},
    )
    assert response.json() == {"text": "hello", "provider": "fake"}
    assert fake_llm.requests[0].json_mode is True


def test_basic_topics(client, fake_llm) -> None:
    fake_llm.answers = ['{"topics": ["a", "b"]}']
    response = client.post("/api/topics/basic", json={"niche": "AI", "duration": "10 minutes"})
    assert response.json() == {"topics": ["a", "b"]}


def test_invalid_ai_output_maps_to_502_with_code(client, fake_llm) -> None:
    fake_llm.answers = ["garbage"]
    response = client.post("/api/topics/unique", json={"niche": "AI"})
    assert response.status_code == 502
    assert response.json()["code"] == "invalid_ai_response"


def test_all_providers_failed_maps_to_503(client, fake_llm) -> None:
    fake_llm.answers = [RuntimeError("down")]
    response = client.post("/api/script", json={
        "topic": "t", "niche": "n", "duration": "1 minutes", "style": "Cinematic",
    })
    assert response.status_code == 503
    assert response.json()["code"] == "all_providers_failed"


def test_split_scenes_accepts_camel_case(client, fake_llm) -> None:
    fake_llm.answers = ['["s1", "s2"]']
    response = client.post("/api/scenes/split", json={"script": "x", "sceneCount": 2})
    assert response.json() == {"scenes": ["s1", "s2"]}


def test_validation_error_shape(client) -> None:
    response = client.post("/api/scenes/split", json={"script": "x", "sceneCount": 0})
    assert response.status_code == 422
    assert response.json()["code"] == "validation_error"


def test_tts_without_key_returns_503(client) -> None:
    response = client.post("/api/tts", json={"text": "hello"})
    assert response.status_code == 503
    assert response.json()["error"] == "UNREAL_SPEECH_API_KEY not configured"


def test_image_falls_back_to_pollinations_url_when_download_fails(client) -> None:
    response = client.post("/api/generate-image", json={"prompt": "a red fox"})
    assert response.json()["url"].startswith("https://image.pollinations.ai/prompt/a%20red%20fox?")


def test_video_endpoints(client) -> None:
    clip = client.post("/api/generate-video", json={"imageUrl": "x", "prompt": "y"}).json()
    assert clip["status"] == "completed"

    final = client.post("/api/compile-video", json={
        "scenes": [{"id": 1, "videoUrl": clip["url"], "text": "ignored"}],
        "audioUrl": "a.mp3",
    }).json()
    assert final["status"] == "success"


def test_script_export_docx(client) -> None:
    response = client.post("/api/script/export", json={
        "topic": "t", "script": "hello", "format": "docx",
    })
    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]


def test_project_export_zip(client) -> None:
    import base64
    import io
    import zipfile

    from openpyxl import load_workbook

    png = "data:image/png;base64," + base64.b64encode(b"\x89PNG fake").decode()
    mp4 = "data:video/mp4;base64," + base64.b64encode(b"fake mp4").decode()
    response = client.post("/api/export/project", json={
        "niche": "AI", "topic": "The AI Revolution!", "script": "Hello.", "summary": "Hi.",
        "style": "Cinematic", "duration": "1 minutes", "aspectRatio": "16:9", "sceneCount": 2,
        "audioUrl": "http://127.0.0.1:1/unreachable.mp3",  # fails -> reported, not fatal
        "finalVideoUrl": mp4,
        "scenes": [
            {"id": 1, "text": "Scene one", "prompt": "p1", "selectedImage": png, "videoUrl": mp4},
            {"id": 2, "text": "Scene two", "prompt": "p2", "selectedImage": png, "videoUrl": mp4},
        ],
    })
    assert response.status_code == 200
    assert 'filename="The_AI_Revolution.zip"' in response.headers["content-disposition"]

    archive = zipfile.ZipFile(io.BytesIO(response.content))
    assert sorted(archive.namelist()) == [
        "final_video.mp4", "images/img1.png", "images/img2.png",
        "project_details.xlsx", "script.txt", "videos/vid1.mp4", "videos/vid2.mp4",
    ]
    assert archive.read("images/img1.png") == b"\x89PNG fake"

    workbook = load_workbook(io.BytesIO(archive.read("project_details.xlsx")))
    assert workbook.sheetnames == ["Overview", "Scenes", "Files"]
    assert workbook["Scenes"]["C2"].value == "p1"
    statuses = {row[0]: row[1] for row in workbook["Files"].iter_rows(min_row=2, values_only=True)}
    assert statuses["audio/narration.mp3"].startswith("missing")


# ---------------------------------------------------------------------- #
# AI selection popup
# ---------------------------------------------------------------------- #
def test_ai_options_marks_providers_without_key_unavailable(client) -> None:
    tasks = {t["id"]: t for t in client.get("/api/ai-options").json()["tasks"]}
    assert list(tasks) == ["topics", "script", "voice", "image", "video", "compile"]
    topics = {o["id"]: o for o in tasks["topics"]["options"]}
    # The test services only have a Groq key.
    assert topics["groq"]["available"] is True
    assert topics["gemini"]["available"] is False
    assert "GEMINI_API_KEY" in topics["gemini"]["reason"]
    # No Ollama / GPU worker in tests: local options are listed but unavailable.
    assert topics["ollama"]["available"] is False and topics["ollama"]["local"] is True
    video = {o["id"]: o for o in tasks["video"]["options"]}
    assert set(video) == {"ltx-video", "svd-xt", "wan-1.3b", "cogvideox-2b"}
    assert all(not o["available"] for o in video.values())
    assert {o["id"] for o in tasks["compile"]["options"]} == {"ffmpeg", "moviepy", "kdenlive"}


def test_tts_voices_per_provider(client) -> None:
    kokoro = client.get("/api/tts/voices", params={"provider": "kokoro"}).json()["voices"]
    assert kokoro[0]["id"] == "af_heart"
    unreal = client.get("/api/tts/voices").json()["voices"]
    assert unreal[0]["id"] == "Sierra"


def test_local_tts_without_worker_returns_503(client) -> None:
    response = client.post("/api/tts", json={"text": "hello", "provider": "kokoro"})
    assert response.status_code == 503
    assert response.json()["code"] == "provider_not_configured"


def test_ollama_provider_id_is_accepted_and_bad_ids_rejected(client) -> None:
    ok = client.post("/api/script/summary", json={"script": "x", "provider": "ollama:qwen3:8b"})
    assert ok.status_code == 502  # resolved, then the (offline) Ollama call fails
    assert ok.json()["code"] == "provider_failed"
    bad = client.post("/api/script/summary", json={"script": "x", "provider": "mystery"})
    assert bad.status_code == 422


def test_verify_reports_invalid_key(client) -> None:
    # The offline HTTP mock answers every request with 404 -> provider error.
    status = client.post("/api/ai-options/verify", json={"provider": "groq"}).json()
    assert status["available"] is False
    assert "Groq" in status["reason"]


def test_image_with_unconfigured_provider_returns_503(client) -> None:
    response = client.post("/api/generate-image", json={"prompt": "a fox", "provider": "gemini"})
    assert response.status_code == 503
    assert response.json()["code"] == "provider_not_configured"
